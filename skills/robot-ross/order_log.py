#!/usr/bin/env python3
"""
order_log.py
------------
Manages the Robot Ross order ledger (orders.xlsx).

Columns:
  Received At | Buyer | Order Type | Content | Produced At | Status | Video URL | Posted | Shipped

Usage (from bob_ross.py):
    import order_log
    row = order_log.start(buyer="Alice", order_type="sketch", content="a cat on a fence")
    order_log.end(row, status="success")

Standalone — update a field manually:
    python3 order_log.py set-video 5 "https://youtube.com/..."
    python3 order_log.py set-posted 5
    python3 order_log.py set-shipped 5
    python3 order_log.py list
"""

import os
import sys
import argparse
from datetime import datetime

ORDERS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "orders.xlsx")

# Column indices (1-based for openpyxl)
COL_RECEIVED   = 1
COL_BUYER      = 2
COL_ORDER_TYPE = 3
COL_CONTENT    = 4
COL_PRODUCED   = 5
COL_STATUS     = 6
COL_VIDEO_URL  = 7
COL_POSTED     = 8
COL_SHIPPED    = 9

HEADERS = [
    "Received At", "Buyer", "Order Type", "Content",
    "Produced At", "Status", "Video URL", "Posted", "Shipped",
]


def _load():
    import openpyxl
    import openpyxl.styles as st

    if os.path.exists(ORDERS_FILE):
        return openpyxl.load_workbook(ORDERS_FILE)

    # Create new workbook with styled header
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(HEADERS)

    bold = st.Font(bold=True)
    fill = st.PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill

    # Column widths
    widths = [20, 15, 12, 40, 20, 12, 50, 10, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb.save(ORDERS_FILE)
    return wb


def start(buyer, order_type, content):
    """
    Append a new in-progress order row.
    Returns the row number (int) to pass to end().
    """
    wb = _load()
    ws = wb.active
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([now, buyer or "", order_type, content, "", "in_progress", "", "", ""])
    wb.save(ORDERS_FILE)
    return ws.max_row


def end(row, status, video_url=None):
    """
    Update the order row with completion time and final status.
    """
    wb = _load()
    ws = wb.active
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.cell(row=row, column=COL_PRODUCED).value = now
    ws.cell(row=row, column=COL_STATUS).value   = status
    if video_url:
        ws.cell(row=row, column=COL_VIDEO_URL).value = video_url
    wb.save(ORDERS_FILE)


def set_video(row, url):
    wb = _load()
    ws = wb.active
    ws.cell(row=row, column=COL_VIDEO_URL).value = url
    wb.save(ORDERS_FILE)
    print(f"Row {row}: video URL set.")


def get_last_video_url():
    """Return the last non-empty Video URL from the orders sheet, or None."""
    wb = _load()
    ws = wb.active
    result = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        val = row[COL_VIDEO_URL - 1]
        if val:
            result = str(val)
    return result


def set_posted(row):
    wb = _load()
    ws = wb.active
    ws.cell(row=row, column=COL_POSTED).value = datetime.now().strftime("%Y-%m-%d")
    wb.save(ORDERS_FILE)
    print(f"Row {row}: marked as posted.")


def set_shipped(row):
    wb = _load()
    ws = wb.active
    ws.cell(row=row, column=COL_SHIPPED).value = datetime.now().strftime("%Y-%m-%d")
    wb.save(ORDERS_FILE)
    print(f"Row {row}: marked as shipped.")


def list_orders():
    wb = _load()
    ws = wb.active
    fmt = "{:<4} {:<20} {:<12} {:<12} {:<20} {:<12} {:<8} {:<8}"
    print(fmt.format("#", "Received At", "Buyer", "Type", "Produced At", "Status", "Posted", "Shipped"))
    print("-" * 100)
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        print(fmt.format(
            i,
            str(row[0] or "")[:20],
            str(row[1] or "")[:12],
            str(row[2] or "")[:12],
            str(row[4] or "")[:20],
            str(row[5] or "")[:12],
            str(row[7] or "")[:8],
            str(row[8] or "")[:8],
        ))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Robot Ross order ledger")
    sub = parser.add_subparsers(dest="cmd")

    sub.add_parser("list", help="Print all orders")

    p_vid = sub.add_parser("set-video", help="Set proof-of-work video URL for a row")
    p_vid.add_argument("row",  type=int)
    p_vid.add_argument("url")

    p_post = sub.add_parser("set-posted", help="Mark a row as posted")
    p_post.add_argument("row", type=int)

    p_ship = sub.add_parser("set-shipped", help="Mark a row as shipped")
    p_ship.add_argument("row", type=int)

    args = parser.parse_args()

    if args.cmd == "list":         list_orders()
    elif args.cmd == "set-video":  set_video(args.row, args.url)
    elif args.cmd == "set-posted": set_posted(args.row)
    elif args.cmd == "set-shipped":set_shipped(args.row)
    else:
        parser.print_help()
